"""El reporte asociado entrega el ZIP o identifica el comprobante sin TC."""

from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.v1.deps import empresa_actual
from app.api.v1.routes import reporte_asociado as ruta
from app.db.database import get_db
from app.domain.comprobante import Libro
from app.services import reporte_asociado

RUC = "20000000000"
PERIODO = "202406"
URL = f"/empresas/{RUC}/periodos/{PERIODO}/reporte-asociado"


def comprobante(libro, tipo_cambio=None, raw_sire=None):
    return {
        "libro": libro.value,
        "origen": "sire",
        "serie_numero": "F001-1",
        "tipo_cp": "01",
        "serie": "F001",
        "numero": "1",
        "tipo_doc_identidad": "6",
        "documento_contraparte": "20000000001",
        "razon_social": "Empresa de prueba",
        "fecha_emision": "2024-06-03",
        "moneda": "USD",
        "tipo_cambio": tipo_cambio,
        "porcentaje_igv": 18,
        "base_imponible": 100,
        "igv": 18,
        "base_imponible_dg": 100,
        "igv_dg": 18,
        "total": 118,
        "glosa": "Servicio de prueba",
        "extra": {"raw_sire": raw_sire} if raw_sire is not None else {},
    }


@pytest.fixture
def contexto(monkeypatch, tmp_path):
    empresa = {"_id": "empresa-de-prueba", "ruc": RUC, "usuario": "PRUEBA"}
    filas = {Libro.COMPRAS: [], Libro.VENTAS: []}
    temporales = []
    carpeta_temporal = tmp_path / "temporales"
    carpeta_temporal.mkdir()

    async def listar(db, empresa_id, periodo, libro, skip, limit):
        assert empresa_id == empresa["_id"]
        assert periodo == PERIODO
        return filas[libro][skip:skip + limit]

    crear_temporal = reporte_asociado.tempfile.NamedTemporaryFile

    def temporal(**kwargs):
        archivo = crear_temporal(dir=carpeta_temporal, **kwargs)
        temporales.append(Path(archivo.name))
        return archivo

    monkeypatch.setattr(ruta.repo_periodos, "obtener", AsyncMock(return_value={"periodo": PERIODO}))
    monkeypatch.setattr(reporte_asociado.comprobantes, "listar", listar)
    monkeypatch.setattr(
        reporte_asociado.pagos_vouchers, "adjuntar_pagos", AsyncMock(return_value=[])
    )
    monkeypatch.setattr(
        reporte_asociado.almacen_pdf.settings, "SUNAT_DATA_DIR", str(tmp_path / "pdfs")
    )
    monkeypatch.setattr(reporte_asociado.tempfile, "NamedTemporaryFile", temporal)

    app = FastAPI()
    app.include_router(ruta.router, prefix="/empresas/{ruc}/periodos/{periodo}")
    app.dependency_overrides[get_db] = lambda: None
    app.dependency_overrides[empresa_actual] = lambda: empresa
    with TestClient(app) as cliente:
        yield SimpleNamespace(cliente=cliente, filas=filas, temporales=temporales)


@pytest.mark.parametrize("libro", [Libro.COMPRAS, Libro.VENTAS])
@pytest.mark.parametrize("tipo_cambio", [None, 0])
def test_sin_tc_devuelve_422_con_el_comprobante_y_borra_el_zip(contexto, libro, tipo_cambio):
    contexto.filas[libro].append(comprobante(libro, tipo_cambio))

    respuesta = contexto.cliente.get(URL)

    assert respuesta.status_code == 422
    detalle = respuesta.json()["detail"]
    assert "sin tipo de cambio SUNAT" in detalle["mensaje"]
    [pendiente] = detalle["comprobantes_sin_tc"]
    assert pendiente["serie_numero"] == "F001-1"
    assert pendiente["moneda"] == "USD"
    assert pendiente["libro"] == libro.value
    assert pendiente["documento_contraparte"] == "20000000001"
    assert pendiente["razon_social"] == "Empresa de prueba"
    assert len(contexto.temporales) == 1
    assert not contexto.temporales[0].exists()


@pytest.mark.parametrize("como_json", [False, True])
def test_el_zip_recupera_tc_de_sunat_y_contiene_ambos_libros_y_pdfs(contexto, como_json):
    raw_sire = {"mtoTipoCambio": "3.740", "codMoneda": "USD"}
    if como_json:
        raw_sire = json.dumps(raw_sire)
    for libro in (Libro.COMPRAS, Libro.VENTAS):
        contexto.filas[libro].append(comprobante(libro, raw_sire=raw_sire))
        pdf = reporte_asociado.almacen_pdf.ruta_pdf(RUC, libro, PERIODO, "01", "F001", "1")
        pdf.parent.mkdir(parents=True)
        pdf.write_bytes(b"%PDF prueba")

    respuesta = contexto.cliente.get(URL)

    assert respuesta.status_code == 200
    assert respuesta.headers["content-type"] == "application/zip"
    assert "PRUEBA-JUNIO2024.zip" in respuesta.headers["content-disposition"]
    with zipfile.ZipFile(io.BytesIO(respuesta.content)) as archivo:
        assert set(archivo.namelist()) == {
            "PRUEBA-JUNIO2024.xlsx",
            "comprobantes compra/",
            "comprobantes compra/facturas/F001-1.pdf",
            "comprobantes venta/",
            "comprobantes venta/facturas/F001-1.pdf",
        }
        libro = load_workbook(io.BytesIO(archivo.read("PRUEBA-JUNIO2024.xlsx")))
        assert libro.sheetnames == [
            "Correlatividad cp", "Registro de ventas", "Registro de compras"
        ]
        assert libro["Registro de compras"]["W4"].value == 3.74
        assert libro["Registro de ventas"]["Q4"].value == 3.74
        assert libro["Registro de compras"]["S4"].value == 118
        assert libro["Registro de ventas"]["P4"].value == 118
    assert len(contexto.temporales) == 1
    assert not contexto.temporales[0].exists()


def test_glosa_pendiente_sigue_bloqueando_la_descarga(contexto):
    fila = comprobante(Libro.COMPRAS, tipo_cambio=3.74)
    fila.pop("glosa")
    contexto.filas[Libro.COMPRAS].append(fila)

    respuesta = contexto.cliente.get(URL)

    assert respuesta.status_code == 409
    assert "Completa las glosas" in respuesta.json()["detail"]
    assert contexto.temporales == []
