"""Alta de empresas: lectura del Excel, registro inmediato y pipeline en la cola."""

from __future__ import annotations

import asyncio
import io
from datetime import UTC, datetime
from unittest.mock import AsyncMock

import pytest
from bson import ObjectId
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.api.v1.routes import empresas as ruta
from app.core.auth import usuario_actual
from app.db.database import get_db
from app.domain.carga_empresas import EstadoFila, FilaCarga, Modalidad, Motivo
from app.domain.jobs import Job, TipoJob
from app.repositories import cargas_empresas as repo_cargas
from app.repositories import empresas as repo_empresas
from app.services import carga_empresas_service as servicio
from app.services.cola.errores import ErrorTransitorio
from app.services.scraping_sunat import CredencialesSolError

RUC_A = "20610202251"
RUC_B = "20603391692"
RUC_C = "20100070970"
CORREO = "contador@example.com"
# Token SIRE con el CIIU 4759 en `userdata.map.ddpData.ddp_ciiu`, sin firma.
TOKEN_CIIU_4759 = (
    "eyJhbGciOiJub25lIn0."
    "eyJ1c2VyZGF0YSI6eyJtYXAiOnsiZGRwRGF0YSI6eyJkZHBfY2lpdSI6IjQ3NTkifX19fQ."
)


def excel(*filas, cabecera=True) -> bytes:
    wb = Workbook()
    hoja = wb.active
    if cabecera:
        hoja.append(["Razón social", "RUC", "Usuario", "Contraseña"])
    for f in filas:
        hoja.append(list(f))
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


# --- Lectura del Excel ------------------------------------------------------


def test_lee_las_cuatro_columnas_y_salta_la_cabecera():
    filas = servicio.leer_excel(excel(("EMPRESA A", RUC_A, "USU", "clave", "ignorada")))
    assert filas == [FilaCarga(2, "EMPRESA A", RUC_A, "USU", "clave")]


def test_sin_cabecera_la_primera_fila_es_una_empresa():
    filas = servicio.leer_excel(excel(("EMPRESA A", RUC_A, "USU", "clave"), cabecera=False))
    assert filas[0].fila == 1


def test_los_numeros_de_excel_se_leen_como_texto_y_las_filas_vacias_se_saltan():
    vacia = (None, None, None, None)
    filas = servicio.leer_excel(excel((" A ", int(RUC_A), "USU", 12345.0), vacia))
    assert [(f.razon_social, f.ruc, f.password) for f in filas] == [("A", RUC_A, "12345")]


@pytest.mark.parametrize("contenido", [b"", b"no es excel", excel()])
def test_rechaza_lo_que_no_sirve(contenido):
    with pytest.raises(servicio.ExcelInvalido):
        servicio.leer_excel(contenido)


def test_tiene_un_tope_de_filas():
    filas = [(f"E{i}", RUC_A, "U", "c") for i in range(201)]
    with pytest.raises(servicio.ExcelInvalido):
        servicio.leer_excel(excel(*filas))


def test_la_plantilla_trae_las_cabeceras():
    hoja = load_workbook(io.BytesIO(servicio.plantilla())).active
    assert [c.value for c in hoja[1]] == ["Razón social", "RUC", "Usuario", "Contraseña"]


# --- Almacén en memoria -----------------------------------------------------


class Almacen:
    def __init__(self) -> None:
        self.empresas: dict[str, dict] = {}
        self.cargas: dict[str, dict] = {}
        self.encolados: list[dict] = []

    async def rucs_existentes(self, db, rucs):
        return {r for r in rucs if r in self.empresas}

    async def obtener_por_ruc(self, db, ruc):
        return self.empresas.get(ruc)

    async def crear_empresa(self, db, datos):
        empresa = {**datos, "_id": ObjectId()}
        self.empresas[datos["ruc"]] = empresa
        return empresa

    async def actualizar(self, db, empresa_id, cambios):
        for empresa in self.empresas.values():
            if empresa["_id"] == empresa_id:
                empresa.update(cambios)
                return empresa
        return None

    async def guardar_token(self, db, empresa_id, token):
        await self.actualizar(db, empresa_id, {"sunat_token": token})

    async def crear_carga(self, db, documento):
        carga_id = str(ObjectId())
        self.cargas[carga_id] = {**documento, "_id": ObjectId(carga_id)}
        return carga_id

    async def obtener_carga(self, db, carga_id):
        return self.cargas.get(carga_id)

    async def listar_cargas(self, db, limit=50):
        return list(self.cargas.values())

    async def guardar_fila(self, db, carga_id, numero, cambios):
        for f in self.cargas[carga_id]["filas"]:
            if f["fila"] == numero:
                f.update(cambios)
        await self.recalcular(db, carga_id)

    async def recalcular(self, db, carga_id):
        carga = self.cargas[carga_id]
        vivas = [f for f in carga["filas"] if f["estado"] == EstadoFila.PENDIENTE.value]
        carga["estado"] = "en_progreso" if vivas else "completada"

    async def encolar(self, db, tipo, ruc, periodo, libro=None, **kwargs):
        self.encolados.append({"tipo": tipo, "ruc": ruc, **kwargs})
        return Job(tipo=tipo, ruc=ruc, periodo=periodo, **kwargs)


@pytest.fixture
def almacen(monkeypatch):
    a = Almacen()
    for nombre, funcion in {
        "rucs_existentes": a.rucs_existentes,
        "obtener_por_ruc": a.obtener_por_ruc,
        "crear": a.crear_empresa,
        "actualizar": a.actualizar,
        "guardar_token_sunat": a.guardar_token,
    }.items():
        monkeypatch.setattr(repo_empresas, nombre, funcion)
    for nombre, funcion in {
        "crear": a.crear_carga,
        "obtener": a.obtener_carga,
        "listar": a.listar_cargas,
        "guardar_fila": a.guardar_fila,
        "recalcular": a.recalcular,
    }.items():
        monkeypatch.setattr(repo_cargas, nombre, funcion)
    monkeypatch.setattr(servicio.cola, "encolar", a.encolar)
    return a


def registrar(filas, modalidad=Modalidad.MASIVA, **kwargs):
    return asyncio.run(
        servicio.registrar(None, filas, modalidad=modalidad, registrado_por=CORREO, **kwargs)
    )


def test_registra_las_validas_y_encola_su_alta(almacen):
    carga = registrar([
        FilaCarga(2, "EMPRESA A", RUC_A, "USU_A", "clave-muy-secreta"),
        FilaCarga(3, "", RUC_B, "USU_B", "otra"),
        FilaCarga(4, "EMPRESA A BIS", RUC_A, "USU_A", "clave"),
    ])

    assert [(f["ruc"], f["estado"], f["motivos"]) for f in carga["filas"]] == [
        (RUC_A, "pendiente", [Motivo.EN_PROCESO.value]),
        (RUC_B, "no_agregada", [Motivo.FALTA_RAZON_SOCIAL.value]),
        (RUC_A, "no_agregada", [Motivo.DUPLICADO_EN_ARCHIVO.value]),
    ]
    empresa = almacen.empresas[RUC_A]
    assert empresa["nombre"] == "EMPRESA A"
    assert empresa["registro"]["modalidad"] == "masiva"
    assert empresa["registro"]["por"] == CORREO
    assert empresa["password"] != "clave-muy-secreta"
    assert "clave-muy-secreta" not in repr(carga)

    [encolado] = almacen.encolados
    assert encolado["tipo"] is TipoJob.ALTA_EMPRESA
    assert encolado["cola"] == RUC_A
    assert encolado["parametros"]["obtener_credenciales"] is True
    # La contraseña nunca va en los parámetros del trabajo.
    assert "clave" not in repr(encolado["parametros"])


def test_un_ruc_ya_registrado_no_se_toca(almacen):
    almacen.empresas[RUC_A] = {"_id": ObjectId(), "ruc": RUC_A, "nombre": "VIEJA"}
    carga = registrar([FilaCarga(2, "NUEVA", RUC_A, "U", "c")])
    assert carga["filas"][0]["motivos"] == [Motivo.RUC_EXISTENTE.value]
    assert almacen.empresas[RUC_A]["nombre"] == "VIEJA"
    assert carga["estado"] == "completada"


def test_en_el_alta_individual_no_se_piden_credenciales_en_la_cola(almacen):
    registrar([FilaCarga(1, "", RUC_A, "U", "c")], modalidad=Modalidad.INDIVIDUAL)
    assert almacen.encolados[0]["parametros"]["obtener_credenciales"] is False


# --- Pipeline de la cola -----------------------------------------------------


@pytest.fixture
def sunat(monkeypatch, almacen):
    """SUNAT que responde bien: token con CIIU y ficha RUC."""

    async def renovar(db, empresa):
        await repo_empresas.guardar_token_sunat(db, empresa["_id"], TOKEN_CIIU_4759)
        return TOKEN_CIIU_4759, None

    async def ficha(db, empresa):
        return await repo_empresas.actualizar(db, empresa["_id"], {
            "actividades_economicas": [
                {"tipo": "PRINCIPAL", "ciiu": "5610", "descripcion": "RESTAURANTES"}
            ],
            "ficha_ruc": {"ruc": empresa["ruc"], "razon_social": "RAZON DE LA FICHA SAC"},
        })

    from app.services import credenciales_sunat_service

    async def credenciales(db, empresa):
        await repo_empresas.actualizar(db, empresa["_id"], {
            "sunat_client_id": "id", "sunat_client_secret": "secreto",
        })
        return {"mensaje": "ok"}

    mocks = {
        "token": AsyncMock(side_effect=renovar),
        "ficha": AsyncMock(side_effect=ficha),
        "credenciales": AsyncMock(side_effect=credenciales),
    }
    monkeypatch.setattr(servicio.auth, "renovar_token", mocks["token"])
    monkeypatch.setattr(servicio.ficha_ruc_service, "actualizar_empresa", mocks["ficha"])
    monkeypatch.setattr(credenciales_sunat_service, "obtener", mocks["credenciales"])
    return mocks


def alta(almacen, fila=2, intentos=1, max_intentos=3, credenciales=True):
    # Individual: la razón social es opcional y así se ve que sale de la ficha.
    carga = registrar([FilaCarga(fila, "", RUC_A, "U", "c")], modalidad=Modalidad.INDIVIDUAL)
    job = Job(
        tipo=TipoJob.ALTA_EMPRESA, ruc=RUC_A, periodo="", intentos=intentos,
        max_intentos=max_intentos,
        parametros={
            "carga_id": str(carga["_id"]), "fila": fila, "obtener_credenciales": credenciales,
        },
    )
    resultado = asyncio.run(servicio.alta_empresa(None, job, AsyncMock()))
    return resultado, almacen.cargas[str(carga["_id"])]


def test_el_alta_completa_credenciales_ciiu_y_rubro(almacen, sunat):
    resultado, carga = alta(almacen)

    empresa = almacen.empresas[RUC_A]
    assert empresa["sunat_client_id"] == "id"
    assert empresa["ciiu"] == "4759"  # el del token manda sobre la ficha
    assert empresa["rubro"] == "Comercio"
    # Sin razón social: se toma de la ficha RUC.
    assert empresa["nombre"] == "RAZON DE LA FICHA SAC"
    assert resultado["estado"] == "agregada"
    assert carga["filas"][0]["estado"] == "agregada"
    assert carga["estado"] == "completada"


def test_sin_token_el_ciiu_sale_de_la_ficha(almacen, sunat):
    sunat["token"].side_effect = None
    sunat["token"].return_value = (None, "401")
    _, carga = alta(almacen, intentos=3, max_intentos=3)
    assert almacen.empresas[RUC_A]["ciiu"] == "5610"
    assert carga["filas"][0]["estado"] == "agregada_con_observaciones"


def test_un_fallo_pasajero_se_reintenta_mientras_queden_intentos(almacen, sunat):
    sunat["token"].side_effect = None
    sunat["token"].return_value = (None, "SUNAT no responde")
    with pytest.raises(ErrorTransitorio):
        alta(almacen, intentos=1, max_intentos=3)
    carga = next(iter(almacen.cargas.values()))
    assert carga["filas"][0]["estado"] == "pendiente"
    assert "reintento" in carga["filas"][0]["motivos"][0]


def test_una_clave_sol_rechazada_no_se_reintenta(almacen, sunat):
    sunat["credenciales"].side_effect = CredencialesSolError("Usuario o Clave Incorrectos")
    _, carga = alta(almacen, intentos=1, max_intentos=3)
    fila = carga["filas"][0]
    assert fila["estado"] == "agregada_con_observaciones"
    assert "clave SOL rechazada" in fila["motivos"][0]
    sunat["token"].assert_not_awaited()
    # La ficha RUC es pública: el CIIU sale de ahí aunque la clave falle.
    assert almacen.empresas[RUC_A]["ciiu"] == "5610"


def test_sin_ciiu_en_ninguna_fuente_queda_con_observaciones(almacen, sunat):
    sunat["token"].side_effect = None
    sunat["token"].return_value = (None, "401")
    sunat["ficha"].side_effect = RuntimeError("Consulta RUC caída")
    _, carga = alta(almacen, intentos=3, max_intentos=3)
    motivos = carga["filas"][0]["motivos"]
    assert any(m.startswith(Motivo.ERROR_SUNAT.value) for m in motivos)
    assert any(m.startswith(Motivo.ERROR_CIIU.value) for m in motivos)
    assert "rubro" not in almacen.empresas[RUC_A]


# --- Rutas ------------------------------------------------------------------


@pytest.fixture
def cliente(almacen):
    ruta.limiter.reset()
    app = FastAPI()
    app.state.limiter = ruta.limiter
    app.include_router(ruta.router, prefix="/empresas")
    app.dependency_overrides[get_db] = lambda: None
    app.dependency_overrides[usuario_actual] = lambda: {"email": CORREO, "rol": "admin"}
    return TestClient(app)


def test_alta_individual_responde_la_empresa_con_su_carga(cliente, almacen):
    r = cliente.post("/empresas", json={"ruc": RUC_A, "usuario": "USU", "password": "clave"})
    assert r.status_code == 201
    cuerpo = r.json()
    assert cuerpo["registro"]["modalidad"] == "individual"
    assert cuerpo["registro"]["por"] == CORREO
    assert cuerpo["carga_id"] in almacen.cargas
    assert "password" not in cuerpo


def test_alta_individual_de_un_ruc_existente_da_409(cliente, almacen):
    almacen.empresas[RUC_A] = {"_id": ObjectId(), "ruc": RUC_A}
    r = cliente.post("/empresas", json={"ruc": RUC_A, "usuario": "USU", "password": "clave"})
    assert r.status_code == 409


def test_alta_individual_rechaza_un_ruc_con_digito_incorrecto(cliente):
    r = cliente.post("/empresas", json={"ruc": "20123456789", "usuario": "U", "password": "c"})
    assert r.status_code == 422


def test_carga_masiva_acepta_el_excel_y_se_consulta_por_id(cliente, almacen):
    r = cliente.post(
        "/empresas/cargas",
        files={"archivo": ("empresas.xlsx", excel(("EMPRESA C", RUC_C, "U", "c")))},
    )
    assert r.status_code == 202
    carga_id = r.json()["carga_id"]

    carga = cliente.get(f"/empresas/cargas/{carga_id}").json()
    assert carga["registrado_por"] == CORREO
    assert carga["archivo"] == "empresas.xlsx"
    assert carga["filas"][0]["ruc"] == RUC_C


def test_carga_masiva_rechaza_lo_que_no_es_excel(cliente):
    assert cliente.post(
        "/empresas/cargas", files={"archivo": ("empresas.csv", b"a,b")}
    ).status_code == 400
    assert cliente.post(
        "/empresas/cargas", files={"archivo": ("empresas.xlsx", b"no es excel")}
    ).status_code == 400


def test_el_reporte_lista_estado_y_motivo(cliente, almacen):
    carga_id = cliente.post(
        "/empresas/cargas",
        files={"archivo": ("e.xlsx", excel(("A", "20123456789", "U", "c")))},
    ).json()["carga_id"]
    almacen.cargas[carga_id]["creado_en"] = datetime.now(UTC)

    r = cliente.get(f"/empresas/cargas/{carga_id}/reporte")

    assert r.status_code == 200
    hoja = load_workbook(io.BytesIO(r.content)).active
    fila = [c.value for c in hoja[2]]
    assert fila[4] == "No agregada"
    assert fila[5] == Motivo.RUC_INVALIDO.value


def test_las_rutas_de_cargas_no_chocan_con_la_de_una_empresa(cliente, almacen):
    # `/cargas/plantilla` y `/cargas` se declaran antes de `/{ruc}`.
    assert cliente.get("/empresas/cargas/plantilla").status_code == 200
    assert cliente.get("/empresas/cargas").status_code == 200
