"""Los vouchers de Apaclla Bot en la plantilla Contasis: van en la fila del
comprobante que pagan, y los que no tienen comprobante, en una hoja aparte."""

from openpyxl import load_workbook

from app.domain.comprobante import Libro
from app.services.comprobante_service import serializar
from app.services.plantilla_excel import excel_plantilla

YAPE = {
    "id": "a", "libro": "ventas", "fuente": "yape", "medio_pago": "Yape",
    "codigo_medio_pago": "003", "nro_operacion": "0012345678", "fecha": "2026-09-19",
    "total": "150.00", "moneda": "PEN", "contraparte": "JUAN PEREZ",
}
PLIN = {**YAPE, "id": "b", "fuente": "plin", "medio_pago": "Plin", "nro_operacion": "5566",
        "total": "80.00"}


def _boleta(**cambios):
    dato = serializar({
        "origen": "sire", "libro": "ventas", "tipo_cp": "03", "serie": "B001", "numero": "45",
        "serie_numero": "B001-45", "moneda": "PEN", "total": 150.0,
    })
    dato.update(cambios)
    return dato


def _columna(hoja, cabecera):
    return next(c.column_letter for c in hoja[2] if c.value == cabecera)


def test_la_fila_de_ventas_lleva_el_medio_de_pago_y_la_operacion():
    wb = load_workbook(excel_plantilla(
        [_boleta(pagos=[YAPE]), _boleta(serie_numero="B001-46", numero="46")], Libro.VENTAS
    ))
    hoja = wb["FORMATO_VENTAS"]

    # MEDIO DE PAGO de Contasis: el código de la Tabla 1 de SUNAT.
    assert hoja["AN4"].value == "003"
    assert hoja[f"{_columna(hoja, 'Medio de pago (voucher)')}4"].value == "Yape"
    assert hoja[f"{_columna(hoja, 'N.º de operación')}4"].value == "0012345678"
    # El comprobante sin voucher queda como antes.
    assert hoja["AN5"].value is None
    assert hoja[f"{_columna(hoja, 'N.º de operación')}5"].value is None
    assert "Vouchers sin comprobante" not in wb.sheetnames


def test_varios_pagos_de_un_mismo_comprobante_se_unen():
    hoja = load_workbook(excel_plantilla([_boleta(pagos=[YAPE, PLIN])], Libro.VENTAS)).active
    assert hoja[f"{_columna(hoja, 'Medio de pago (voucher)')}4"].value == "Yape / Plin"
    assert hoja[f"{_columna(hoja, 'N.º de operación')}4"].value == "0012345678 / 5566"


def test_compras_no_tiene_columna_de_medio_de_pago_de_contasis():
    compra = _boleta(libro="compras", tipo_cp="01", serie="F001", pagos=[YAPE])
    wb = load_workbook(excel_plantilla([compra], Libro.COMPRAS))
    hoja = wb.worksheets[0]

    assert hoja["AN4"].value is None
    assert hoja[f"{_columna(hoja, 'N.º de operación')}4"].value == "0012345678"


def test_los_vouchers_sin_comprobante_van_en_su_hoja_y_no_en_el_registro():
    wb = load_workbook(excel_plantilla([_boleta()], Libro.VENTAS, vouchers_sin_comprobante=[PLIN]))

    registro = wb["FORMATO_VENTAS"]
    assert registro["G5"].value == "TOTAL"  # una sola fila de datos
    hoja = wb["Vouchers sin comprobante"]
    assert [c.value for c in hoja[1]] == [
        "Medio de pago", "N.º de operación", "Fecha", "Contraparte", "Total", "Moneda"
    ]
    assert hoja["A2"].value == "Plin"
    assert hoja["B2"].value == "5566"
    assert float(hoja["E2"].value) == 80.0
