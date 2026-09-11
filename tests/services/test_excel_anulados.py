import pytest
from openpyxl import load_workbook

from app.domain.comprobante import Libro
from app.services.comprobante_service import serializar
from app.services.plantilla_excel import excel_plantilla


@pytest.mark.parametrize("libro", [Libro.COMPRAS, Libro.VENTAS])
def test_excel_incluye_anulados_originales_sunat(libro):
    datos = [serializar({
        "origen": "sire", "serie_numero": "F001-1", "glosa": "Editada",
        "detalle_sunat": [{"descripcion": "Anulado"}],
    }), serializar({"origen": "sire", "serie_numero": "F001-2", "glosa": "Anulado"})]
    wb = load_workbook(excel_plantilla(datos, libro))
    assert wb['Anulados'].max_row == 2
    assert wb['Anulados']['A2'].value == 'F001-1'
    assert wb['Anulados']['E2'].value == 'Anulado'
    hoja = wb.worksheets[0]
    columna = next(c.column_letter for c in hoja[2] if c.value == 'Observación')
    assert hoja[f'{columna}4'].value is None
    assert hoja[f'{columna}5'].value is None
