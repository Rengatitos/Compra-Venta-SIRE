import pytest
from openpyxl import load_workbook

from app.domain.comprobante import Libro
from app.services.comprobante_service import serializar
from app.services.plantilla_excel import excel_plantilla
from app.services.revision_comprobantes import incompleto


def test_solo_avisa_despues_de_consultar_y_excluye_anulados():
    assert not incompleto({})
    assert incompleto({'glosa_consultada': True, 'razon_social': 'Cliente', 'documento_contraparte': '-'})
    assert not incompleto({'glosa_consultada': True, 'razon_social': 'Cliente', 'documento_contraparte': '00123456'})
    assert not incompleto({'origen': 'sire', 'detalle_sunat': [{'descripcion': 'Anulado'}]})


def test_edicion_manual_completa_campos_sin_cambiar_sire():
    original = {'razon_social': '-', 'documento_contraparte': '', 'contraparte_manual': {
        'razon_social': 'Cliente', 'documento_contraparte': '00123456'}}
    salida = serializar(original)
    assert salida['razon_social'] == 'Cliente'
    assert salida['documento_contraparte'] == '00123456'
    assert original['razon_social'] == '-'


@pytest.mark.parametrize('libro', [Libro.COMPRAS, Libro.VENTAS])
def test_excel_incluye_incompletos_excluye_anulados_del_registro(libro):
    filas = [serializar({'origen': 'sire', 'total': 100, 'moneda': 'PEN',
        'detalle_sunat': [{'descripcion': 'Anulado'}]}),
        serializar({'origen': 'sire', 'total': 60, 'moneda': 'PEN', 'glosa_consultada': True})]
    wb = load_workbook(excel_plantilla(filas, libro))
    col = 'S' if libro == Libro.COMPRAS else 'P'
    assert wb.worksheets[0][f'{col}4'].value == 60
    assert wb.worksheets[0][f'{col}5'].value == f'=SUM({col}4:{col}4)'
    assert wb['Anulados'].max_row == 2
