from __future__ import annotations

import csv
import json
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader

from app.services.clasificador.hashing import stable_id
from app.services.clasificador.text import compact_value, join_nonempty, split_text

SUPPORTED_EXTENSIONS = {".xlsx", ".csv", ".json", ".jsonl", ".md", ".txt", ".pdf"}


@dataclass(slots=True)
class Chunk:
    id: str
    content: str
    source: str
    metadata: dict[str, Any]
    source_weight: float = 1.0

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> Chunk:
        return cls(**data)


class DocumentLoader:
    def __init__(self, knowledge_dir: Path):
        self.knowledge_dir = knowledge_dir

    def list_documents(self) -> list[Path]:
        paths: list[Path] = []
        for path in self.knowledge_dir.rglob("*"):
            relative_parts = {part.lower() for part in path.relative_to(self.knowledge_dir).parts}
            if relative_parts & {"historico", "histórico", "historical", "history"}:
                continue
            if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS and not path.name.startswith("."):
                paths.append(path)
        return sorted(paths)

    def load(self, path: Path, source_weight: float = 1.0) -> list[Chunk]:
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            return self._load_xlsx(path, source_weight)
        if suffix == ".csv":
            return self._load_csv(path, source_weight)
        if suffix in {".json", ".jsonl"}:
            return self._load_json(path, source_weight)
        if suffix in {".md", ".txt"}:
            return self._load_text(path, source_weight)
        if suffix == ".pdf":
            return self._load_pdf(path, source_weight)
        return []

    def _relative(self, path: Path) -> str:
        try:
            return path.resolve().relative_to(self.knowledge_dir.resolve()).as_posix()
        except Exception:
            return path.name

    def _load_xlsx(self, path: Path, source_weight: float) -> list[Chunk]:
        source = self._relative(path)
        wb = load_workbook(path, read_only=True, data_only=True)
        chunks: list[Chunk] = []
        try:
            for ws in wb.worksheets:
                rows = ws.iter_rows(values_only=True)
                headers: list[str] | None = None
                header_row = 0
                category_groups: dict[str, list[str]] = {}
                parent_groups: dict[str, list[str]] = {}
                for row_no, row in enumerate(rows, start=1):
                    values = [compact_value(v) for v in row]
                    if not any(values):
                        continue
                    if headers is None:
                        headers = [v if v else f"col_{i+1}" for i, v in enumerate(values)]
                        header_row = row_no
                        continue
                    data = {headers[i]: values[i] for i in range(min(len(headers), len(values))) if values[i] != ""}
                    if not data:
                        continue
                    lower_map = {k.strip().lower(): v for k, v in data.items()}
                    text_key = next((k for k in data if k.strip().lower() in {"texto_rag", "text_rag", "contenido_rag"}), None)
                    if text_key:
                        main = data[text_key]
                        extras = [f"{k}: {v}" for k, v in data.items() if k != text_key and k.lower() in {
                            "cuenta", "descripcion", "categoria", "criterio_inclusion", "ruc", "razon_social",
                            "ciiu", "ciiu_v4", "tipo", "codigo", "libro", "reporte", "centro_costos",
                            "proveedor_ruc", "contraparte_ruc", "numero_documento", "cuenta_base",
                            "cuenta_total", "glosa", "descripcion_operacion", "tipo_cp"
                        }]
                        content = join_nonempty([main, join_nonempty(extras, "; ")], "\n")
                    else:
                        content = "\n".join(f"{k}: {v}" for k, v in data.items())
                    meta = {
                        "kind": "xlsx_row",
                        "sheet": ws.title,
                        "row": row_no,
                        "header_row": header_row,
                    }
                    for key, value in data.items():
                        lk = key.strip().lower()
                        if lk in {
                            "cuenta", "codigo", "ruc", "ciiu", "ciiu_v4", "tipo", "libro", "reporte",
                            "centro_costos", "categoria", "proveedor_ruc", "contraparte_ruc",
                            "numero_documento", "cuenta_base", "cuenta_total", "glosa",
                            "descripcion_operacion", "tipo_cp", "cuenta_padre",
                            "es_cuenta_hoja", "cuenta_raiz"
                        }:
                            meta[lk] = value
                    chunks.append(Chunk(
                        id=stable_id(source, ws.title, str(row_no), content[:200]),
                        content=content,
                        source=source,
                        metadata=meta,
                        source_weight=source_weight,
                    ))

                    # Para planes/catálogos jerárquicos crea chunks compactos de contexto.
                    # No agrega conocimiento nuevo: solo agrupa filas vecinas para que un modelo
                    # pequeño vea alternativas de una misma rama contable en un único Top-K.
                    cuenta = lower_map.get("cuenta")
                    descripcion = lower_map.get("descripcion")
                    if cuenta and descripcion:
                        compact = f"Cuenta {cuenta}: {descripcion}"
                        if lower_map.get("centro_costos"):
                            compact += f". Centro de costos: {lower_map['centro_costos']}"
                        if lower_map.get("relacion_reporte"):
                            compact += f". Relación: {lower_map['relacion_reporte']}"
                        categoria = lower_map.get("categoria")
                        if categoria:
                            category_groups.setdefault(categoria, []).append(compact)
                        parent = lower_map.get("cuenta_padre")
                        parent_desc = lower_map.get("descripcion_padre")
                        if parent:
                            key = f"{parent} - {parent_desc or ''}".strip()
                            parent_groups.setdefault(key, []).append(compact)

                for group_kind, groups in (("xlsx_category_group", category_groups), ("xlsx_parent_group", parent_groups)):
                    for group_name, lines in groups.items():
                        if len(lines) < 2:
                            continue
                        group_text = f"Grupo contable: {group_name}\n" + "\n".join(lines)
                        for part_i, part in enumerate(split_text(group_text, max_chars=1800, overlap_chars=120)):
                            chunks.append(Chunk(
                                id=stable_id(source, ws.title, group_kind, group_name, str(part_i)),
                                content=part,
                                source=source,
                                metadata={"kind": group_kind, "sheet": ws.title, "group": group_name, "part": part_i},
                                source_weight=source_weight,
                            ))
        finally:
            wb.close()
        return chunks

    def _load_csv(self, path: Path, source_weight: float) -> list[Chunk]:
        source = self._relative(path)
        chunks: list[Chunk] = []
        with path.open("r", encoding="utf-8-sig", errors="ignore", newline="") as f:
            reader = csv.DictReader(f)
            for row_no, row in enumerate(reader, start=2):
                data = {str(k): compact_value(v) for k, v in row.items() if k and compact_value(v)}
                if not data:
                    continue
                content = "\n".join(f"{k}: {v}" for k, v in data.items())
                chunks.append(Chunk(stable_id(source, str(row_no), content[:200]), content, source, {"kind":"csv_row","row":row_no}, source_weight))
        return chunks

    def _load_json(self, path: Path, source_weight: float) -> list[Chunk]:
        source = self._relative(path)
        records: list[Any] = []
        if path.suffix.lower() == ".jsonl":
            with path.open("r", encoding="utf-8", errors="ignore") as f:
                for line in f:
                    line = line.strip()
                    if line:
                        try:
                            records.append(json.loads(line))
                        except json.JSONDecodeError:
                            continue
        else:
            with path.open("r", encoding="utf-8", errors="ignore") as f:
                data = json.load(f)
            records = data if isinstance(data, list) else [data]
        chunks: list[Chunk] = []
        for i, record in enumerate(records):
            content = json.dumps(record, ensure_ascii=False, separators=(",", ":"))
            chunks.append(Chunk(stable_id(source, str(i), content[:200]), content, source, {"kind":"json_record","index":i}, source_weight))
        return chunks

    def _load_text(self, path: Path, source_weight: float) -> list[Chunk]:
        source = self._relative(path)
        text = path.read_text(encoding="utf-8", errors="ignore")
        return [
            Chunk(stable_id(source, str(i), chunk[:200]), chunk, source, {"kind":"text","chunk":i}, source_weight)
            for i, chunk in enumerate(split_text(text))
        ]

    def _load_pdf(self, path: Path, source_weight: float) -> list[Chunk]:
        source = self._relative(path)
        reader = PdfReader(str(path))
        chunks: list[Chunk] = []
        for page_no, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""
            for local_i, chunk in enumerate(split_text(text, max_chars=1500, overlap_chars=180)):
                chunks.append(Chunk(
                    id=stable_id(source, str(page_no), str(local_i), chunk[:200]),
                    content=chunk,
                    source=source,
                    metadata={"kind":"pdf","page":page_no,"chunk":local_i},
                    source_weight=source_weight,
                ))
        return chunks
