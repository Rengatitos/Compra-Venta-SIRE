"""Libro conjunto y control de numeración por registro, emisor, tipo y serie."""
import io
import re
from collections import defaultdict, Counter
from copy import copy

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from app.domain.comprobante import Libro
from app.repositories import comprobantes
from app.services.comprobante_service import serializar_lote
from app.services.glosa import obtener_glosa
from app.services.plantilla_excel import excel_plantilla

MESES = ('ENERO FEBRERO MARZO ABRIL MAYO JUNIO JULIO AGOSTO SEPTIEMBRE OCTUBRE NOVIEMBRE DICIEMBRE').split()


def nombre(usuario, periodo):
    return f"{re.sub(r'[^A-Za-z0-9_-]', '', usuario)}-{MESES[int(periodo[4:6])-1]}{periodo[:4]}"


async def registros(db, empresa, periodo):
    salida = {}
    for libro in (Libro.VENTAS, Libro.COMPRAS):
        filas = []
        while True:
            pagina = await comprobantes.listar(db, str(empresa['_id']), periodo, libro=libro, skip=len(filas), limit=500)
            filas.extend(pagina)
            if len(pagina) < 500:
                break
        salida[libro] = filas
    return salida


def estado(registros):
    filas = [f for grupo in registros.values() for f in grupo]
    pendientes = sum(not bool(obtener_glosa(f)) for f in filas)
    return {'habilitado': bool(filas) and pendientes == 0, 'pendientes': pendientes}


def excel(registros, periodo, ruc):
    wb = Workbook()
    control = wb.active
    control.title = 'Correlatividad cp'
    control.append(['CONTROL DE NUMERACIÓN DE COMPROBANTES · RVIE Y RCE'])
    control.append(['Los saltos corresponden solo a los documentos del periodo; no prueban omisiones del emisor.'])
    control.append(['Mes', 'Registro', 'Emisor / Proveedor', 'Tipo', 'Serie', 'Del', 'Al', 'Cantidad', 'Saltos', 'Duplicados', 'Fuera de secuencia por fecha'])
    for libro, filas in registros.items():
        grupos = defaultdict(list)
        for f in filas:
            emisor = ruc if libro == Libro.VENTAS else f.get('documento_contraparte', '')
            grupos[(emisor, f.get('tipo_cp', ''), f.get('serie', ''))].append(f)
        for (emisor, tipo, serie), grupo in sorted(grupos.items(), key=lambda par: str(par[0])):
            numeros = [int(str(f.get('numero'))) for f in grupo if str(f.get('numero', '')).isdigit()]
            orden = sorted(set(numeros))
            saltos = [str(a+1) if b == a+2 else f'{a+1}-{b-1}' for a, b in zip(orden, orden[1:]) if b > a+1]
            duplicados = [str(n) for n, cantidad in Counter(numeros).items() if cantidad > 1]
            por_fecha = sorted(grupo, key=lambda f: (str(f.get('fecha_emision') or ''), int(str(f.get('numero'))) if str(f.get('numero', '')).isdigit() else -1))
            secuencia = [int(str(f['numero'])) for f in por_fecha if str(f.get('numero', '')).isdigit()]
            control.append([MESES[int(periodo[4:6])-1], libro.value, emisor, tipo, serie,
                min(orden) if orden else None, max(orden) if orden else None, len(grupo),
                ', '.join(saltos) or ('Numeración no numérica' if len(numeros) != len(grupo) else ''),
                ', '.join(duplicados), 'Revisar' if any(b < a for a, b in zip(secuencia, secuencia[1:])) else ''])
    control.freeze_panes = 'A4'
    control.auto_filter.ref = f'A3:K{max(3, control.max_row)}'
    for celda in control[3]:
        celda.font = Font(bold=True)
    for col in 'ABCDEFGHIJK':
        control.column_dimensions[col].width = 24
    for libro in (Libro.VENTAS, Libro.COMPRAS):
        original = load_workbook(excel_plantilla(serializar_lote(registros[libro]), libro)).worksheets[0]
        hoja = wb.create_sheet('Registro de ventas' if libro == Libro.VENTAS else 'Registro de compras')
        for fila in original:
            for celda in fila:
                destino = hoja.cell(celda.row, celda.column, celda.value)
                destino._style = copy(celda._style)
        for clave, dimension in original.column_dimensions.items():
            hoja.column_dimensions[clave] = copy(dimension)
        for clave, dimension in original.row_dimensions.items():
            hoja.row_dimensions[clave] = copy(dimension)
        for rango in original.merged_cells.ranges:
            hoja.merge_cells(str(rango))
        hoja.freeze_panes = original.freeze_panes
    for fila in control:
        for celda in fila:
            if isinstance(celda.value, str):
                celda.data_type = 's'
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
